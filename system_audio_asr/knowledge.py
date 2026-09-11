"""知识库：多条可启用的资料条目，拼入 AI 系统提示词供答题时优先引用。

存储独立于 config.json（``knowledge.json``）：C# Overlay 会整体重写
config.json，知识库若混写其中会在浮窗保存设置时被抹掉，因此沿用
phone_share.json 的独立文件约定。

条目结构：``{"id": str, "title": str, "content": str, "enabled": bool}``。
支持从 PDF / DOCX / TXT / MD 提取文本后作为条目内容写入（不落盘原文件）。
"""
from __future__ import annotations

import io
import json
import os
import re
import secrets
import threading
from pathlib import Path
from typing import Any

APP_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")) / "WasapiParaformerOverlay"
KNOWLEDGE_PATH = APP_DIR / "knowledge.json"

MAX_ENTRIES = 200
MAX_TITLE_CHARS = 80
MAX_CONTENT_CHARS = 200000
# 拼进提示词时的单条与总长度上限：防止长知识库撑爆模型上下文窗口
MAX_CONTEXT_ENTRY_CHARS = 6000
MAX_CONTEXT_TOTAL_CHARS = 24000

# 支持上传解析的扩展名
DOCUMENT_SUFFIXES = {".txt", ".md", ".markdown", ".docx", ".pdf"}

_store_lock = threading.Lock()


def _normalize_entry(raw: Any) -> dict[str, Any] | None:
    if not isinstance(raw, dict):
        return None
    raw_title = raw.get("title")
    raw_content = raw.get("content")
    # 只接受标量：列表/字典等容器类型视为非法输入，避免 str(list) 产生垃圾内容
    if isinstance(raw_content, (list, dict, tuple, set)):
        return None
    title = ("" if raw_title is None else str(raw_title)).strip()[:MAX_TITLE_CHARS]
    content = ("" if raw_content is None else str(raw_content)).strip()
    if not content:
        return None
    if len(content) > MAX_CONTENT_CHARS:
        content = content[:MAX_CONTENT_CHARS]
    entry_id = str(raw.get("id") or "").strip() or secrets.token_hex(8)
    return {
        "id": entry_id,
        "title": title or "未命名条目",
        "content": content,
        "enabled": bool(raw.get("enabled", True)),
    }


def load_entries(path: Path | None = None) -> list[dict[str, Any]]:
    """读取全部知识库条目；文件缺失或损坏时返回空列表（不抛异常，避免拖垮服务启动）。"""
    target = path or KNOWLEDGE_PATH
    with _store_lock:
        try:
            raw = json.loads(target.read_text(encoding="utf-8-sig"))
        except (OSError, ValueError):
            return []
    items = raw.get("entries") if isinstance(raw, dict) else raw
    if not isinstance(items, list):
        return []
    entries: list[dict[str, Any]] = []
    for item in items[:MAX_ENTRIES]:
        normalized = _normalize_entry(item)
        if normalized:
            entries.append(normalized)
    return entries


def save_entries(entries: list[dict[str, Any]], path: Path | None = None) -> list[dict[str, Any]]:
    """原子写入：先写 .tmp 再 os.replace，避免并发读时读到半截文件。"""
    target = path or KNOWLEDGE_PATH
    normalized: list[dict[str, Any]] = []
    for item in (entries or [])[:MAX_ENTRIES]:
        entry = _normalize_entry(item)
        if entry:
            normalized.append(entry)
    with _store_lock:
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_suffix(".tmp")
        temporary.write_text(
            json.dumps({"entries": normalized}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        os.replace(temporary, target)
    return normalized


def upsert_entry(
    entry: dict[str, Any], path: Path | None = None
) -> dict[str, Any]:
    """新增或按 id 更新一条；返回写入后的条目。"""
    entries = load_entries(path)
    entry_id = str((entry or {}).get("id") or "").strip()
    candidate = dict(entry or {})
    candidate["id"] = entry_id or secrets.token_hex(8)
    if entry_id:
        for index, existing in enumerate(entries):
            if existing["id"] == entry_id:
                entries[index] = {**existing, **candidate}
                break
        else:
            entries.append(candidate)
    else:
        entries.append(candidate)
    saved = save_entries(entries, path)
    target_id = candidate["id"]
    for item in saved:
        if item["id"] == target_id:
            return item
    return saved[-1] if saved else {}


def delete_entry(entry_id: str, path: Path | None = None) -> list[dict[str, Any]]:
    entries = [item for item in load_entries(path) if item["id"] != entry_id]
    return save_entries(entries, path)


def clear_entries(path: Path | None = None) -> list[dict[str, Any]]:
    return save_entries([], path)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _parse_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 依赖，无法解析 DOCX") from exc
    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs]
    # 表格里的内容（简历常用表格排版）也要取出来
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(part for part in parts if part)


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("缺少 pypdf 依赖，无法解析 PDF") from exc
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            parts.append(text.strip())
    return "\n".join(parts)


def parse_document(data: bytes, filename: str) -> str:
    """从上传文档提取纯文本。

    仅支持带文本层的文件；扫描件/图片版 PDF 提取结果为空，会抛出明确提示，
    由调用方引导用户手动粘贴（不在此处接入 OCR）。
    """
    if not data:
        raise RuntimeError("文件内容为空")
    suffix = Path(str(filename or "")).suffix.lower()
    if suffix not in DOCUMENT_SUFFIXES:
        raise RuntimeError("不支持的文件格式（仅支持 PDF / DOCX / TXT / MD）")
    # 解析库对损坏文件会抛 BadZipFile / PdfStreamError 等原生异常，
    # 统一包成 RuntimeError，让接口层返回 400 提示而不是 500。
    try:
        if suffix == ".docx":
            text = _parse_docx(data)
        elif suffix == ".pdf":
            text = _parse_pdf(data)
        else:
            text = _decode_text(data)
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError(
            "文件解析失败，可能已损坏或不是有效的 "
            + suffix.lstrip(".").upper()
            + " 文件"
        ) from exc
    text = _clean_text(text)
    if not text.strip():
        raise RuntimeError("未能提取到文字，可能是扫描件/图片格式，请手动粘贴内容")
    return text[:MAX_CONTENT_CHARS]


def _clean_text(text: str) -> str:
    """压缩连续空行与多余空白，保留段落换行。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def build_context_block(
    entries: list[dict[str, Any]] | None = None,
    *,
    limit: int = MAX_CONTEXT_TOTAL_CHARS,
) -> str:
    """把启用的条目拼成提示词片段；超出总量上限时按条目顺序截断。"""
    items = load_entries() if entries is None else entries
    sections: list[str] = []
    used = 0
    for entry in items:
        if not entry.get("enabled", True):
            continue
        content = str(entry.get("content") or "").strip()
        if not content:
            continue
        if len(content) > MAX_CONTEXT_ENTRY_CHARS:
            content = content[:MAX_CONTEXT_ENTRY_CHARS].rstrip() + "…"
        block = "[" + str(entry.get("title") or "知识库") + "]\n" + content
        if used + len(block) > limit:
            break
        sections.append(block)
        used += len(block)
    if not sections:
        return ""
    return "\n\n".join(sections) + "\n\n"


def public_entries(path: Path | None = None) -> list[dict[str, Any]]:
    """供前端展示：附上每条的字数，便于用户判断是否过长。"""
    return [
        {**entry, "chars": len(entry["content"])}
        for entry in load_entries(path)
    ]